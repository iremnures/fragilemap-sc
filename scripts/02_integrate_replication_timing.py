import csv
import openpyxl
from pathlib import Path
from collections import defaultdict, Counter

# ============================================================
# Paths
# ============================================================

PROJECT_DIR = Path.home() / "fragilemap_sc"

BREAK_FILE = (
    PROJECT_DIR
    / "data"
    / "processed"
    / "human_break_regions_enriched.csv"
)

RT_WORKBOOK = (
    PROJECT_DIR
    / "data"
    / "raw"
    / "41467_2026_76451_MOESM4_ESM.xlsx"
)

OUTPUT_RT = (
    PROJECT_DIR
    / "data"
    / "processed"
    / "human_break_regions_with_rt.csv"
)

OUTPUT_CLASSIFIED = (
    PROJECT_DIR
    / "data"
    / "processed"
    / "human_break_regions_rt_classified.csv"
)

# ============================================================
# Cell-line mapping
# ============================================================

SHEET_MAP = {
    "U2OS": "U2OS",
    "hTERT-RPE1": "RPE1",
    "BJ-hTERT": "BJ"
}


def normalize_chrom(value):
    """
    Normalize chromosome labels to chr-style notation.
    """
    value = str(value).strip()

    if value.lower().startswith("chr"):
        return "chr" + value[3:]

    return "chr" + value


def is_number(value):
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


# ============================================================
# Detect Pseudobulk RT untreated column
# ============================================================

def detect_rt_column(ws):
    """
    Supplementary Data 2 uses multi-row headers.

    Search the first several rows and combine header text
    vertically for each column. The target column must contain
    both 'Pseudobulk RT' and 'Untreated'.
    """

    header_rows = []

    for row in ws.iter_rows(
        min_row=1,
        max_row=10,
        values_only=True
    ):
        header_rows.append(row)

    max_cols = max(len(r) for r in header_rows)

    for col_idx in range(max_cols):

        pieces = []

        for row in header_rows:

            if col_idx < len(row):
                value = row[col_idx]

                if value is not None:
                    pieces.append(str(value))

        combined = " ".join(pieces).lower()

        if (
            "pseudobulk" in combined
            and "rt" in combined
            and "untreated" in combined
        ):
            return col_idx

    raise RuntimeError(
        f"Could not detect untreated Pseudobulk RT column "
        f"in sheet: {ws.title}"
    )


# ============================================================
# Load break-region dataset
# ============================================================

break_regions = []

with open(BREAK_FILE, newline="") as infile:

    reader = csv.DictReader(infile)

    original_fields = reader.fieldnames

    for row in reader:

        row["start"] = int(row["start"])
        row["end"] = int(row["end"])
        row["frequency"] = int(row["frequency"])

        row["chrom"] = normalize_chrom(row["chrom"])

        break_regions.append(row)

assert len(break_regions) == 644, (
    f"Expected 644 break regions, found {len(break_regions)}"
)

# ============================================================
# Read only RT information required from Supplementary Data 2
# ============================================================

print("\nREPLICATION-TIMING INTEGRATION")
print("=" * 70)

wb = openpyxl.load_workbook(
    RT_WORKBOOK,
    read_only=True,
    data_only=True
)

# Structure:
# rt_data[cell_type][chrom] =
#     [(start, end, rt_value), ...]

rt_data = {
    "U2OS": defaultdict(list),
    "RPE1": defaultdict(list),
    "BJ": defaultdict(list)
}

for sheet_name, cell_type in SHEET_MAP.items():

    ws = wb[sheet_name]

    rt_col = detect_rt_column(ws)

    n_windows = 0
    n_values = 0

    for row in ws.iter_rows(values_only=True):

        # Need chrom, start, end + RT column
        if len(row) <= rt_col:
            continue

        chrom = row[0]
        start = row[1]
        end = row[2]
        rt_value = row[rt_col]

        # Detect actual genomic-data rows rather than
        # relying on a fixed header row number.
        if chrom is None:
            continue

        chrom_text = str(chrom).strip()

        if not chrom_text.lower().startswith("chr"):
            continue

        if not is_number(start) or not is_number(end):
            continue

        start = int(float(start))
        end = int(float(end))

        n_windows += 1

        if rt_value is None or not is_number(rt_value):
            continue

        rt_value = float(rt_value)

        chrom_text = normalize_chrom(chrom_text)

        rt_data[cell_type][chrom_text].append(
            (
                start,
                end,
                rt_value
            )
        )

        n_values += 1

    # Coordinate-sort within chromosome
    for chrom in rt_data[cell_type]:
        rt_data[cell_type][chrom].sort(
            key=lambda x: x[0]
        )

    print(
        f"{cell_type:<5} | "
        f"RT column={rt_col + 1:<2} | "
        f"genomic windows={n_windows:<6} | "
        f"non-missing RT={n_values}"
    )

wb.close()

# ============================================================
# Calculate overlap-weighted RT for every break region
# ============================================================

def weighted_rt(cell_type, chrom, break_start, break_end):

    bins = rt_data[cell_type].get(chrom, [])

    weighted_sum = 0.0
    total_overlap = 0

    for bin_start, bin_end, rt_value in bins:

        # Sorted intervals allow early exit
        if bin_start >= break_end:
            break

        if bin_end <= break_start:
            continue

        overlap_start = max(
            break_start,
            bin_start
        )

        overlap_end = min(
            break_end,
            bin_end
        )

        overlap_bp = (
            overlap_end
            - overlap_start
        )

        if overlap_bp > 0:

            weighted_sum += (
                rt_value
                * overlap_bp
            )

            total_overlap += overlap_bp

    if total_overlap == 0:
        return None

    return weighted_sum / total_overlap


def classify_rt(value):

    if value > 0.25:
        return "Early"

    if value < -0.25:
        return "Late"

    return "Mid"


integrated = []

missing_rt = 0

for row in break_regions:

    rt_value = weighted_rt(
        row["cell_type"],
        row["chrom"],
        row["start"],
        row["end"]
    )

    if rt_value is None:
        missing_rt += 1
        rt_class = ""
    else:
        rt_class = classify_rt(rt_value)

    new_row = dict(row)

    new_row[
        "mean_pseudobulk_rt_untreated"
    ] = rt_value

    new_row["rt_class"] = rt_class

    integrated.append(new_row)

# ============================================================
# Write RT-enriched table
# ============================================================

rt_fields = (
    list(original_fields)
    + ["mean_pseudobulk_rt_untreated"]
)

with open(
    OUTPUT_RT,
    "w",
    newline=""
) as outfile:

    writer = csv.DictWriter(
        outfile,
        fieldnames=rt_fields
    )

    writer.writeheader()

    for row in integrated:

        output_row = {
            key: row[key]
            for key in rt_fields
        }

        writer.writerow(output_row)

# ============================================================
# Write classified table
# ============================================================

classified_fields = (
    rt_fields
    + ["rt_class"]
)

with open(
    OUTPUT_CLASSIFIED,
    "w",
    newline=""
) as outfile:

    writer = csv.DictWriter(
        outfile,
        fieldnames=classified_fields
    )

    writer.writeheader()

    for row in integrated:

        output_row = {
            key: row[key]
            for key in classified_fields
        }

        writer.writerow(output_row)

# ============================================================
# Validation
# ============================================================

assert len(integrated) == 644, (
    f"Expected 644 integrated regions; "
    f"found {len(integrated)}"
)

assert missing_rt == 0, (
    f"{missing_rt} break regions have no RT value"
)

expected_counts = {
    "U2OS": {
        "Early": 64,
        "Mid": 68,
        "Late": 63
    },
    "RPE1": {
        "Early": 98,
        "Mid": 100,
        "Late": 86
    },
    "BJ": {
        "Early": 37,
        "Mid": 79,
        "Late": 49
    }
}

print("\nRT CLASS VALIDATION")
print("=" * 70)

for cell_type in [
    "U2OS",
    "RPE1",
    "BJ"
]:

    observed = Counter(
        row["rt_class"]
        for row in integrated
        if row["cell_type"] == cell_type
    )

    print(
        f"{cell_type:<5} | "
        f"Early={observed['Early']:<3} | "
        f"Mid={observed['Mid']:<3} | "
        f"Late={observed['Late']:<3}"
    )

    assert observed == Counter(
        expected_counts[cell_type]
    ), (
        f"{cell_type}: RT-class counts "
        f"do not match validated results"
    )

print("\nValidation: PASS")
print("Integrated break regions:", len(integrated))
print("Missing RT values:", missing_rt)

print("\nSaved:")
print(OUTPUT_RT)
print(OUTPUT_CLASSIFIED)
